#!/usr/bin/env python3
"""Build the PowerPoint deck for The Mad Hour from the same content as installationpresentation.html.

    python3 tools/build-pptx.py

Writes deck/the-mad-hour.pptx: 16:9, one agenda point per slide, the home page's
generated hero scene as the stage, a smooth fade between slides on click.

The agenda lives in installationpresentation.html and is read from there, so the two decks can never
drift apart. Text is real text, not a picture, so names can be fixed in PowerPoint
on the day; it wants Cormorant Garamond and Jost installed (fonts/install/).
"""

import json, re, sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
BUILD = ROOT / "deck"
TMP = BUILD / "assets"

# the palette, straight off the page
INK        = RGBColor(0x1E, 0x33, 0x24)
GOLD_INK   = RGBColor(0x7E, 0x5F, 0x3A)
GOLD_LINE  = RGBColor(0xDC, 0xC3, 0xA3)
SAGE_DEEP  = RGBColor(0x7F, 0x96, 0x73)
PAPER      = (251, 240, 227)

DISPLAY, SANS = "Cormorant Garamond", "Jost"

W, H = Inches(13.333), Inches(7.5)          # 16:9
RULE_Y      = Inches(7.5 * 0.537)            # where the hairline sits, as on the page
RULE_W      = Inches(1.05)
SIDE        = Inches(1.5)
CREST_W     = Inches(0.69)
CREST_TOP   = Inches(0.34)

# topic point size and measure, one step per length, mirroring .topic-xl/lg/md/sm.
# The measure is the CSS max-width in characters, turned into inches at roughly
# 0.44 em per character of Cormorant, so lines break where they break on the page.
STEPS = ((18, 52, 15), (40, 42, 18), (62, 34, 22), (10**6, 27, 26))
HONORIFIC = re.compile(r"^(Leo Lion|MJF Lion|Lion|Leo)\s+")


def js_text(s):
    """A JavaScript string literal as text: \\u escapes resolved, real UTF-8 left alone."""
    return s.encode("latin-1", "backslashreplace").decode("unicode_escape").replace("\\'", "'")


def read_agenda():
    """Pull EVENT and AGENDA out of installationpresentation.html so there is one source of truth."""
    src = (ROOT / "installationpresentation.html").read_text(encoding="utf-8")
    def field(name):
        return js_text(re.search(name + r":'((?:[^'\\]|\\.)*)'", src).group(1))
    event = {k: field(k) for k in ("title", "sub", "when", "where")}
    body = re.search(r"var AGENDA=\[(.*?)\n  \];", src, re.S).group(1)
    agenda = []
    for row in re.finditer(r"\{ topic:'((?:[^'\\]|\\.)*)', speaker:'((?:[^'\\]|\\.)*)', mins:(\d+) \}", body):
        topic, speaker, mins = row.groups()
        agenda.append({"topic": js_text(topic), "speaker": js_text(speaker), "mins": int(mins)})
    if len(agenda) < 2:
        sys.exit("could not read AGENDA out of installationpresentation.html")
    return event, agenda


def stage_plate(px=(1920, 1080)):
    """The hero render, cropped as CSS crops it and washed with the same paper
    gradients, baked into one image so the slide text stays editable."""
    TMP.mkdir(parents=True, exist_ok=True)
    out = TMP / "stage.jpg"
    src = Image.open(ROOT / "images" / "hero-wide.webp").convert("RGB")
    tw, th = px
    # background-size: cover, background-position: center 28%
    scale = max(tw / src.width, th / src.height)
    rw, rh = round(src.width * scale), round(src.height * scale)
    img = src.resize((rw, rh), Image.LANCZOS)
    left = (rw - tw) // 2
    top = round((rh - th) * 0.28)
    img = img.crop((left, top, left + tw, top + th))

    # radial-gradient(56% 56% at 50% 48%, paper .78 -> 0 at 84%)
    wash = Image.new("L", (tw, th), 0)
    d = ImageDraw.Draw(wash)
    steps = 160
    rx, ry = 0.56 * tw, 0.56 * th
    cx, cy = 0.50 * tw, 0.48 * th
    for i in range(steps, 0, -1):
        t = i / steps                     # 0 centre -> 1 edge of the gradient box
        a = 0.78 * max(0.0, 1 - t / 0.84) if t < 0.84 else 0.0
        d.ellipse([cx - rx * t, cy - ry * t, cx + rx * t, cy + ry * t], fill=round(a * 255))
    img = Image.composite(Image.new("RGB", (tw, th), PAPER), img, wash)

    # linear-gradient(180deg, paper .18 0%, transparent 30%)
    top_fade = Image.new("L", (tw, th), 0)
    dt = ImageDraw.Draw(top_fade)
    band = round(th * 0.30)
    for y in range(band):
        dt.line([(0, y), (tw, y)], fill=round(0.18 * (1 - y / band) * 255))
    img = Image.composite(Image.new("RGB", (tw, th), PAPER), img, top_fade)

    img.save(out, quality=92, subsampling=0, optimize=True)
    return out


def crest_png():
    out = TMP / "crest.png"
    Image.open(ROOT / "images" / "logo-340.webp").convert("RGBA").save(out)
    return out


def track(run, em):
    """Letter-spacing: OOXML spells it `spc`, in hundredths of a point."""
    run.font._rPr.set("spc", str(round(em * run.font.size.pt * 100)))


def fade_on_click(slide, ms=700):
    """A slow cross-fade, advancing only when the presenter clicks. Written as
    AlternateContent so PowerPoint 2010+ gets the timed fade and everything else
    (Keynote, Google Slides, older PowerPoint) still gets a plain fade."""
    xml = (
        '<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
        ' xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        '<mc:Choice xmlns:p14="http://schemas.microsoft.com/office/powerpoint/2010/main" Requires="p14">'
        f'<p:transition spd="slow" p14:dur="{ms}" advClick="1"><p:fade/></p:transition>'
        '</mc:Choice><mc:Fallback>'
        '<p:transition spd="slow" advClick="1"><p:fade/></p:transition>'
        '</mc:Fallback></mc:AlternateContent>')
    from pptx.oxml import parse_xml
    slide._element.append(parse_xml(xml))


def textbox(slide, y, height, anchor, width=None):
    width = width or (W - 2 * SIDE)
    box = slide.shapes.add_textbox(int((W - width) / 2), y, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def add_stage(slide, plate, crest, crest_w=CREST_W, crest_top=CREST_TOP):
    slide.shapes.add_picture(str(plate), 0, 0, W, H)
    if crest_w:
        ratio = 322 / 340
        slide.shapes.add_picture(str(crest), int((W - crest_w) / 2), crest_top,
                                 crest_w, int(crest_w * ratio))


def add_rule(slide, y, width=RULE_W):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int((W - width) / 2), y, width, Emu(9525))
    line.fill.solid()
    line.fill.fore_color.rgb = GOLD_LINE
    line.line.fill.background()
    line.shadow.inherit = False
    return line


def agenda_slide(prs, plate, crest, item, n, total):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_stage(slide, plate, crest)

    size, measure = next((pt, ch) for limit, pt, ch in STEPS if len(item["topic"]) <= limit)
    tf = textbox(slide, Inches(0.9), RULE_Y - Inches(0.9) - Inches(0.26), MSO_ANCHOR.BOTTOM,
                 width=Inches(measure * 0.44 * size / 72))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 1.06
    run = p.add_run()
    run.text = item["topic"]
    run.font.name, run.font.size, run.font.color.rgb = DISPLAY, Pt(size), INK

    add_rule(slide, RULE_Y)

    tf = textbox(slide, RULE_Y + Inches(0.26), Inches(1.4), MSO_ANCHOR.TOP, width=Inches(8.5))
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 1.4
    for i, whole in enumerate(item["speaker"].split(" / ")):
        if i:
            sep = p.add_run()
            sep.text = "   /   "
            sep.font.name, sep.font.size, sep.font.color.rgb = SANS, Pt(18), GOLD_LINE
        m = HONORIFIC.match(whole)
        if m:
            hon = p.add_run()
            hon.text = m.group(1).upper() + " "
            hon.font.name, hon.font.size, hon.font.color.rgb = SANS, Pt(14.5), GOLD_INK
            track(hon, 0.2)
        name = p.add_run()
        name.text = whole[m.end():] if m else whole
        name.font.name, name.font.size, name.font.color.rgb = SANS, Pt(18), INK
        name.font.bold = False
        track(name, 0.045)

    slide.notes_slide.notes_text_frame.text = (
        f"Point {n} of {total} · {item['mins']} min · {item['speaker']}")
    fade_on_click(slide)
    return slide


def title_slide(prs, plate, crest, head, sub, meta):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_stage(slide, plate, crest, crest_w=Inches(1.2), crest_top=Inches(1.55))

    tf = textbox(slide, Inches(2.75), Inches(1.5), MSO_ANCHOR.TOP)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = head
    run.font.name, run.font.size, run.font.color.rgb = DISPLAY, Pt(60), INK
    # no tracking on the title cards: at display size PowerPoint keeps Cormorant's
    # "Th" ligature and then spaces around it, which splits "Thank you"


    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(6)
    run = p2.add_run()
    run.text = sub
    run.font.name, run.font.size, run.font.color.rgb = DISPLAY, Pt(20), SAGE_DEEP
    run.font.italic = True

    add_rule(slide, Inches(4.5), width=Inches(1.25))

    tf = textbox(slide, Inches(4.76), Inches(1.1), MSO_ANCHOR.TOP)
    for i, line in enumerate(meta):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.space_before = Pt(0 if i == 0 else 4)
        run = p.add_run()
        run.text = line
        run.font.name, run.font.size, run.font.color.rgb = SANS, Pt(11), GOLD_INK
        track(run, 0.03)

    fade_on_click(slide)
    return slide


def main():
    event, agenda = read_agenda()
    BUILD.mkdir(parents=True, exist_ok=True)
    plate, crest = stage_plate(), crest_png()

    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H

    title_slide(prs, plate, crest, event["title"], event["sub"], [event["when"], event["where"]])
    for i, item in enumerate(agenda, 1):
        agenda_slide(prs, plate, crest, item, i, len(agenda))
    title_slide(prs, plate, crest, "Thank you", "Built on bonds.",
                ["Leo Club of Aurelian · District 3231 A4"])

    out = BUILD / "the-mad-hour.pptx"
    prs.save(out)
    print(f"{out.relative_to(ROOT)}: {len(prs.slides.__iter__.__self__._sldIdLst)} slides, "
          f"{out.stat().st_size // 1024} KB")


if __name__ == "__main__":
    main()
